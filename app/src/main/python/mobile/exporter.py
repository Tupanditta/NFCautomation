import os
import json
import logging
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from mobile.attendance import get_attendance_summary, get_calendar_config, get_schedule

def generate_attendance_report(config_dir, cache_dir, period_type="MONTH", semester="q1"):
    """
    Genera un informe detallado en Excel (.xlsx) con 3 hojas.
    """
    try:
        now = datetime.now()
        
        # 1. Resolver semestre si es AUTO
        if semester == "AUTO":
            calendar = get_calendar_config(config_dir)
            from mobile.attendance import get_semester_for_date
            semester = get_semester_for_date(now, calendar) or "q1"

        # 2. Obtener datos según el periodo
        data = get_attendance_summary(config_dir, period_type, semester, "ALL")
        records = data.get("records", [])
        stats = data.get("stats", {})
        
        wb = Workbook()
        
        # --- HOJA 1: RESUMEN GENERAL ---
        ws1 = wb.active
        ws1.title = "Resumen General"
        
        title_font = Font(name='Calibri', size=16, bold=True, color="1F4E78")
        header_font = Font(name='Calibri', bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        stripe_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_side = Side(style='thin', color="BFBFBF")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        def get_perc_fill(value_str):
            try:
                val = float(str(value_str).replace('%', ''))
                if val >= 80: return PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                if val >= 50: return PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                return PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            except: return None

        # Título combinado
        ws1.merge_cells("A1:G1")
        ws1["A1"] = "INFORME DE ASISTENCIA ACADÉMICA"
        ws1["A1"].font = title_font
        ws1["A1"].alignment = center_align
        
        ws1["A2"] = f"Periodo: {period_type} | Cuatrimestre: {semester.upper()}"
        ws1["A3"] = f"Generado el: {now.strftime('%d/%m/%Y %H:%M:%S')}"
        
        # Estadísticas Globales
        ws1.append([])
        ws1.append(["ESTADÍSTICAS GLOBALES"])
        ws1.append(["Total Asistidas", stats.get("total_attended", 0)])
        ws1.append(["Total Faltadas", stats.get("total_missed", 0)])
        ws1.append(["Total Pendientes", stats.get("total_upcoming", 0)])
        ws1.append(["Tasa de Asistencia", f"{stats.get('percentage', 0)}%"])

        # Resumen por Asignatura
        ws1.append([])
        ws1.append(["RESUMEN POR ASIGNATURA"])
        headers = ["Asignatura", "Tipo", "Aula", "Piso", "Edificio", "Total Clases", "Asistidas", "Faltadas", "% Asistencia"]
        ws1.append(headers)
        
        # Agrupar y obtener aula/piso/edificio
        subj_data = {}
        for r in records:
            if r.get("status") in ["HOLIDAY", "WEEKEND", "FREE"]: continue
            s = r["subject"]
            t = r.get("type", "THEORY")
            key = (s, t)
            if key not in subj_data:
                subj_data[key] = {
                    "total": 0, "att": 0, "miss": 0,
                    "rooms": set(), "floors": set(), "buildings": set()
                }
            subj_data[key]["total"] += 1
            if r["status"] == "ATTENDED": subj_data[key]["att"] += 1
            elif r["status"] == "MISSED": subj_data[key]["miss"] += 1

            rm = r.get("room")
            if rm and rm != "-": subj_data[key]["rooms"].add(rm)
            fl = r.get("floor")
            if fl and fl != "-": subj_data[key]["floors"].add(fl)
            bg = r.get("building")
            if bg and bg != "-": subj_data[key]["buildings"].add(bg)

        row_count = 0
        for (s, t), info in sorted(subj_data.items()):
            total_fin = info["att"] + info["miss"]
            perc = f"{round((info['att'] / total_fin * 100), 1) if total_fin > 0 else 0}%"
            room = " / ".join(sorted(info["rooms"])) if info["rooms"] else "-"
            floor = " / ".join(sorted(info["floors"])) if info["floors"] else "-"
            building = " / ".join(sorted(info["buildings"])) if info["buildings"] else "-"
            type_label = "TEORÍA" if t == "THEORY" else "PRÁCTICA"
            
            ws1.append([s, type_label, room, floor, building, info["total"], info["att"], info["miss"], perc])
            row_count += 1
            curr_row = ws1.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws1.cell(row=curr_row, column=col_idx)
                cell.border = thin_border
                if row_count % 2 == 0: cell.fill = stripe_fill
                cell.alignment = center_align if col_idx > 1 else left_align
                if col_idx == 9: cell.fill = get_perc_fill(perc)

        # --- HOJA 2 & 3 ---
        ws2 = wb.create_sheet("Faltas")
        ws2.sheet_properties.tabColor = "FF0000"
        ws2.append(["FECHA", "DÍA", "HORA", "ASIGNATURA", "TIPO"])
        for r in records:
            if r.get("status") == "MISSED":
                t_label = "TEORÍA" if r.get("type") == "THEORY" else "PRÁCTICA"
                ws2.append([r["date"], r["day"].upper(), r["start"], r["subject"].upper(), t_label])

        ws3 = wb.create_sheet("Asistencias")
        ws3.sheet_properties.tabColor = "00B050"
        ws3.append(["FECHA", "DÍA", "HORA", "ASIGNATURA", "TIPO", "MÉTODO"])
        for r in records:
            if r.get("status") == "ATTENDED":
                t_label = "TEORÍA" if r.get("type") == "THEORY" else "PRÁCTICA"
                ws3.append([r["date"], r["day"].upper(), r["start"], r["subject"].upper(), t_label, "NFC/MANUAL"])

        # Estilo final y auto-ajuste (Extremadamente seguro)
        for ws in wb.worksheets:
            # 1. Aplicar estilos básicos a cabeceras de hojas 2 y 3 (si no es la principal)
            if ws.title != "Resumen General":
                for cell in ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = center_align
                    cell.border = thin_border

            # 2. Ajustar anchos ignorando celdas combinadas (MergedCells)
            for col_idx in range(1, ws.max_column + 1):
                max_len = 5 # Mínimo inicial
                column_letter = get_column_letter(col_idx)
                
                for row_idx in range(1, ws.max_row + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    
                    # Comprobación de seguridad para evitar MergedCell
                    if type(cell).__name__ == 'Cell':
                        try:
                            if cell.value is not None:
                                val_str = str(cell.value)
                                if len(val_str) > max_len:
                                    max_len = len(val_str)
                        except:
                            continue
                
                # Aplicar ancho con un pequeño margen
                ws.column_dimensions[column_letter].width = min(max_len + 4, 45)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Informe_Asistencia_{period_type}_{timestamp}.xlsx"
        file_path = os.path.join(cache_dir, filename)
        wb.save(file_path)
        return file_path

    except Exception as e:
        logging.error(f"Error en generador Excel: {str(e)}")
        raise e
